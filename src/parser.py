
from dataclasses import asdict
from src.models import TextBlock
from src.models import ResumeShape
from src.models import ResumeContent
from pathlib import Path
import json
from typing import Any

# pyrefly: ignore [missing-import]
from pptx import Presentation


def parse_resume(resume_path: Path) -> ResumeContent:
    prs = Presentation(resume_path)
    resume_content = ResumeContent()
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                resume_shape = ResumeShape(name=shape.name)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            text_block = TextBlock(
                                text=run.text,
                                bold=run.font.bold or False,
                                font_size=(run.font.size or 0) / 12700,
                            )
                            resume_shape.blocks.append(text_block)
                if resume_shape.blocks:
                    resume_content.shapes.append(resume_shape)
    return resume_content


def save_resume(resume_content: ResumeContent, output_path: Path) -> None:
    """
    Saves the resume content to a JSON file.
    """
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(asdict(resume_content), f, indent=4, ensure_ascii=False)


def extract_resume_sections(resume_path: Path) -> dict[str, Any]:
    """
    Extracts structured, section-level content from the resume presentation template.
    Returns dictionary with experience, projects, education coursework, military bullets,
    and technical tools/concepts.
    """
    resume_path = resume_path.resolve()
    if not resume_path.exists():
        raise FileNotFoundError(f"Resume template not found: {resume_path}")

    prs = Presentation(resume_path)
    slide = prs.slides[0]
    data: dict[str, Any] = {}

    for shape in slide.shapes:
        if shape.name == "TextBox 36" and shape.has_text_frame:
            data["experience_bullets"] = [
                p.text.strip() for p in shape.text_frame.paragraphs[2:] if p.text.strip()
            ]
        elif shape.name == "TextBox 43" and shape.has_text_frame:
            projects = []
            ps = shape.text_frame.paragraphs
            for i in range(0, len(ps), 2):
                if i + 1 < len(ps):
                    projects.append({
                        "title_line": ps[i].text.strip(),
                        "description": ps[i + 1].text.strip(),
                    })
            data["academic_projects"] = projects
        elif shape.name == "TextBox 34" and shape.has_text_frame:
            ps = shape.text_frame.paragraphs
            if len(ps) >= 3:
                data["coursework"] = ps[2].text.strip()
        elif shape.name == "TextBox 38" and shape.has_text_frame:
            data["military_bullets"] = [
                p.text.strip() for p in shape.text_frame.paragraphs[1:] if p.text.strip()
            ]
        elif shape.name == "TextBox 50" and shape.has_text_frame:
            ps = shape.text_frame.paragraphs
            data["tools_raw"] = [ps[i].text.strip() for i in range(2, 6) if i < len(ps)]
            data["concepts_raw"] = [ps[i].text.strip() for i in range(8, 10) if i < len(ps)]
        elif shape.name == "TextBox 17" and shape.has_text_frame:
            data["spoken_languages"] = [
                p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
            ]
        elif shape.name == "TextBox 40" and shape.has_text_frame:
            data["sport_bullets"] = [
                p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
            ]
        elif shape.name == "TextBox 47" and shape.has_text_frame:
            data["volunteering_bullets"] = [
                p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
            ]
        elif shape.name == "Group 51" and hasattr(shape, "shapes"):
            langs = []
            for g in shape.shapes:
                if hasattr(g, "shapes"):
                    for sub in g.shapes:
                        if sub.has_text_frame and sub.text_frame.text.strip():
                            langs.append(sub.text_frame.text.strip())
            if langs:
                data["programming_languages"] = langs

    return data