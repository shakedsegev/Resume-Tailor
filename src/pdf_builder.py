import os
import subprocess
from pathlib import Path
from typing import Optional, Any
# pyrefly: ignore [missing-import]
from pptx import Presentation
# pyrefly: ignore [missing-import]
from pptx.shapes.autoshape import Shape
# pyrefly: ignore [missing-import]
from pptx.util import Pt


def convert_pptx_to_pdf(pptx_path: Path, pdf_path: Path) -> Path:
    """
    Converts a .pptx presentation to a high-resolution .pdf on macOS.
    Uses Apple Keynote via AppleScript, with fallback to Microsoft PowerPoint.
    """
    pptx_path = pptx_path.resolve()
    pdf_path = pdf_path.resolve()

    if not pptx_path.exists():
        raise FileNotFoundError(f"Source presentation not found: {pptx_path}")

    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Try Keynote (Standard on macOS, very fast & clean vector PDF export)
    keynote_script = f'''
    tell application "Keynote"
        set theDoc to open POSIX file "{pptx_path}"
        export theDoc to POSIX file "{pdf_path}" as PDF
        close theDoc saving no
    end tell
    '''
    res = subprocess.run(["osascript", "-e", keynote_script], capture_output=True, text=True)

    if res.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 0:
        return pdf_path

    # 2. Fallback to Microsoft PowerPoint if Keynote is unavailable or encountered an issue
    powerpoint_script = f'''
    tell application "Microsoft PowerPoint"
        set theDoc to open POSIX file "{pptx_path}"
        save theDoc in POSIX file "{pdf_path}" as save as PDF
        close theDoc saving no
    end tell
    '''
    res_ppt = subprocess.run(["osascript", "-e", powerpoint_script], capture_output=True, text=True)

    if res_ppt.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 0:
        return pdf_path

    raise RuntimeError(
        f"PDF conversion failed.\n"
        f"Keynote error: {res.stderr.strip()}\n"
        f"PowerPoint error: {res_ppt.stderr.strip()}"
    )


def update_paragraphs_in_shape(
    shape: Shape,
    new_texts: list[str],
    start_paragraph_idx: int = 0,
) -> None:
    """
    Replaces paragraph text in a PowerPoint text shape while strictly preserving
    original font name, size, bolding, colors, and paragraph styles.
    """
    if not shape.has_text_frame:
        return

    tf = shape.text_frame
    paragraphs = tf.paragraphs

    # Extract template style reference from first available run in the target range
    template_run = None
    for p in paragraphs[start_paragraph_idx:]:
        if p.runs:
            template_run = p.runs[0]
            break

    for i, new_text in enumerate(new_texts):
        target_idx = start_paragraph_idx + i
        if target_idx < len(paragraphs):
            p = paragraphs[target_idx]
            if p.runs:
                # Keep the first run's formatting, update text, clear any leftover runs
                p.runs[0].text = new_text
                for leftover_run in p.runs[1:]:
                    leftover_run.text = ""
            else:
                run = p.add_run()
                run.text = new_text
                _copy_run_style(template_run, run)
        else:
            # If there are more bullets than original paragraphs, add paragraph
            new_p = tf.add_paragraph()
            run = new_p.add_run()
            run.text = new_text
            _copy_run_style(template_run, run)


def _copy_run_style(source_run: Optional[Any], target_run: Any) -> None:
    """Copies font properties (name, size, bold, color) from source to target run."""
    if not source_run:
        return
    try:
        if source_run.font.name:
            target_run.font.name = source_run.font.name
        if source_run.font.size:
            target_run.font.size = source_run.font.size
        if source_run.font.bold is not None:
            target_run.font.bold = source_run.font.bold
        if source_run.font.color and source_run.font.color.type is not None:
            try:
                target_run.font.color.rgb = source_run.font.color.rgb
            except Exception:
                pass
    except Exception:
        pass


def update_projects_in_shape(
    shape: Shape,
    tailored_projects: list[Any],
) -> None:
    """
    Updates Academic Projects in a text shape (e.g. TextBox 43).
    Each project is represented by 2 paragraphs:
      - Title Paragraph: Title (Run 0, bold) + Tech Tags (Run 1, normal)
      - Description Paragraph: Project Description (Run 0, normal)
    """
    if not shape.has_text_frame:
        return

    paragraphs = shape.text_frame.paragraphs
    for i, proj in enumerate(tailored_projects):
        title_p_idx = i * 2
        desc_p_idx = i * 2 + 1

        title_suffix = (
            proj.title_suffix
            if hasattr(proj, "title_suffix")
            else proj.get("title_suffix", "")
        )
        description = (
            proj.description
            if hasattr(proj, "description")
            else proj.get("description", "")
        )

        # Update Title Tech Tags (Run 1)
        if title_p_idx < len(paragraphs):
            p = paragraphs[title_p_idx]
            if len(p.runs) >= 2:
                p.runs[1].text = (
                    title_suffix
                    if title_suffix.startswith(" ")
                    else f" {title_suffix}"
                )
            elif len(p.runs) == 1:
                run = p.add_run()
                run.text = title_suffix
                _copy_run_style(p.runs[0], run)
                run.font.bold = False

        # Update Description Paragraph
        if desc_p_idx < len(paragraphs):
            p = paragraphs[desc_p_idx]
            if p.runs:
                p.runs[0].text = description
                for r in p.runs[1:]:
                    r.text = ""


def inject_tailored_content(
    template_pptx_path: Path,
    output_pptx_path: Path,
    shape_updates: dict[str, dict[str, Any]],
) -> Path:
    """
    Injects tailored text into specified shapes in a PowerPoint template.

    shape_updates format:
    {
        "TextBox 36": {"texts": ["bullet 1", "bullet 2"], "start_idx": 2},
        "TextBox 43": {"projects": [TailoredProject(...), ...]},
    }
    """
    template_pptx_path = template_pptx_path.resolve()
    output_pptx_path = output_pptx_path.resolve()

    if not template_pptx_path.exists():
        raise FileNotFoundError(f"Template PPTX not found: {template_pptx_path}")

    prs = Presentation(template_pptx_path)
    output_pptx_path.parent.mkdir(parents=True, exist_ok=True)

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.name in shape_updates:
                spec = shape_updates[shape.name]
                if "projects" in spec:
                    update_projects_in_shape(shape, spec["projects"])
                else:
                    texts = spec.get("texts", [])
                    start_idx = spec.get("start_idx", 0)
                    update_paragraphs_in_shape(shape, texts, start_paragraph_idx=start_idx)

    prs.save(output_pptx_path)
    return output_pptx_path


def build_tailored_resume(
    template_pptx: Path,
    output_dir: Path,
    shape_updates: dict[str, dict[str, Any]],
    base_name: str = "tailored_resume",
) -> tuple[Path, Path]:
    """
    Complete Step 4 workflow:
    1. Injects tailored content into the PPTX.
    2. Exports high-resolution PDF.
    Returns (pptx_path, pdf_path).
    """
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    pptx_out = output_dir / f"{base_name}.pptx"
    pdf_out = output_dir / f"{base_name}.pdf"

    inject_tailored_content(template_pptx, pptx_out, shape_updates)
    convert_pptx_to_pdf(pptx_out, pdf_out)

    return pptx_out, pdf_out


def reorder_pill_groups(group_shape: Shape, ordered_languages: list[str]) -> None:
    """
    Physically relocates the entire native pill groups in Group 51 row by row,
    preserving each pill's custom oval graphics and avoiding any text wrapping.
    """
    name_to_group = {}
    for g in group_shape.shapes:
        for sub in g.shapes:
            if sub.has_text_frame and sub.text_frame.text.strip():
                name_to_group[sub.text_frame.text.strip()] = g

    # Standard 3-row layout coordinates from Canva
    row_tops = [0, 456352, 912704]
    col_gap = 40000

    # Ensure all 7 languages are accounted for
    all_langs = list(name_to_group.keys())
    full_ordered = [l for l in ordered_languages if l in name_to_group]
    for l in all_langs:
        if l not in full_ordered:
            full_ordered.append(l)

    # Pack into 3 rows: 3 in row 1, 2 in row 2, 2 in row 3
    row1 = full_ordered[:3]
    row2 = full_ordered[3:5]
    row3 = full_ordered[5:]

    for row_idx, row_langs in enumerate([row1, row2, row3]):
        curr_left = 0
        t = row_tops[row_idx]
        for lang in row_langs:
            if lang in name_to_group:
                g = name_to_group[lang]
                g.left = curr_left
                g.top = t
                curr_left += g.width + col_gap


def update_tools_and_concepts(
    shape: Shape,
    tools_lines: list[str],
    core_concepts_lines: list[str],
) -> None:
    """
    Updates TextBox 50:
    - Paragraphs 2..5 (index 2..5): 4 lines of Tools & Platforms
    - Paragraphs 8..9 (index 8..9): 2 lines of Core Concepts
    """
    if not shape.has_text_frame:
        return
    paragraphs = shape.text_frame.paragraphs

    # Update tools lines (P3..P6 -> indices 2..5)
    for i, line in enumerate(tools_lines[:4]):
        idx = 2 + i
        if idx < len(paragraphs) and paragraphs[idx].runs:
            paragraphs[idx].runs[0].text = line

    # Update concepts lines (P9..P10 -> indices 8..9)
    for i, line in enumerate(core_concepts_lines[:2]):
        idx = 8 + i
        if idx < len(paragraphs) and paragraphs[idx].runs:
            paragraphs[idx].runs[0].text = line


def update_coursework(shape: Shape, coursework_line: str) -> None:
    """
    Updates TextBox 34 (Education), paragraph 2 (index 2: Core CS Coursework).
    """
    if not shape.has_text_frame:
        return
    paragraphs = shape.text_frame.paragraphs
    if len(paragraphs) >= 3 and paragraphs[2].runs:
        paragraphs[2].runs[0].text = coursework_line


def update_military_bullets(shape: Shape, military_bullets: list[str]) -> None:
    """
    Updates TextBox 38 (Military Service), paragraphs 1..3 (index 1..3: Bullets).
    """
    if not shape.has_text_frame:
        return
    paragraphs = shape.text_frame.paragraphs
    for i, b in enumerate(military_bullets[:3]):
        idx = 1 + i
        if idx < len(paragraphs) and paragraphs[idx].runs:
            paragraphs[idx].runs[0].text = b


def apply_full_tailoring(prs: Presentation, tailored: Any) -> None:
    """
    Applies comprehensive tailoring across all sections of the presentation.
    """
    slide = prs.slides[0]

    for shape in slide.shapes:
        # 1. Experience (TextBox 36)
        if shape.name == "TextBox 36" and hasattr(tailored, "experience_bullets"):
            update_paragraphs_in_shape(
                shape, tailored.experience_bullets, start_paragraph_idx=2
            )
            for p in shape.text_frame.paragraphs[2:]:
                for r in p.runs:
                    r.font.size = Pt(8.6)

        # 2. Academic Projects (TextBox 43)
        elif shape.name == "TextBox 43" and hasattr(tailored, "academic_projects"):
            update_projects_in_shape(shape, tailored.academic_projects)
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.6)

        # 3. Programming Languages Pills (Group 51)
        elif shape.name == "Group 51" and hasattr(tailored, "ordered_languages") and tailored.ordered_languages:
            reorder_pill_groups(shape, tailored.ordered_languages)

        # 4. Tools & Core Concepts (TextBox 50)
        elif shape.name == "TextBox 50":
            tools = getattr(tailored, "tools_lines", [])
            concepts = getattr(tailored, "core_concepts_lines", [])
            if tools or concepts:
                update_tools_and_concepts(shape, tools, concepts)

        # 5. Education Coursework (TextBox 34)
        elif shape.name == "TextBox 34" and hasattr(tailored, "coursework_line"):
            update_coursework(shape, tailored.coursework_line)

        # 6. Military Leadership (TextBox 38)
        elif shape.name == "TextBox 38" and hasattr(tailored, "military_bullets"):
            update_military_bullets(shape, tailored.military_bullets)
            for p in shape.text_frame.paragraphs[1:]:
                for r in p.runs:
                    r.font.size = Pt(8.6)


def build_full_tailored_resume(
    template_pptx: Path,
    output_dir: Path,
    tailored: Any,
    base_name: str = "tailored_resume",
) -> tuple[Path, Path]:
    """
    Complete end-to-end full resume tailoring:
    1. Loads template presentation.
    2. Applies full-spectrum updates across all sections.
    3. Exports pristine vector PDF.
    Returns (pptx_path, pdf_path).
    """
    template_pptx = template_pptx.resolve()
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    pptx_out = output_dir / f"{base_name}.pptx"
    pdf_out = output_dir / f"{base_name}.pdf"

    prs = Presentation(template_pptx)
    apply_full_tailoring(prs, tailored)
    prs.save(pptx_out)

    convert_pptx_to_pdf(pptx_out, pdf_out)
    return pptx_out, pdf_out
