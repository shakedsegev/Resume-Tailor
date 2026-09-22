from pathlib import Path
from pptx import Presentation
import pytest
from src.pdf_builder import reorder_pill_groups, apply_full_tailoring
from src.ai_engine import TailoredProject

TEMPLATE_PATH = Path("template.pptx")


class DummyTailored:
    def __init__(self):
        self.ordered_languages = ["Python", "C++", "SQL", "Java", "TypeScript", "C", "Assembly x86"]
        self.academic_projects = [
            TailoredProject(
                title_suffix=" | Java, Sockets, Reactor",
                description="Test STOMP description fitting exactly within limits.",
            ),
            TailoredProject(
                title_suffix=" | Java, Multithreading",
                description="Test Linear Algebra description fitting within limits.",
            ),
            TailoredProject(
                title_suffix=" | C++, Memory Management",
                description="Test DJ Station description fitting within limits.",
            ),
        ]
        self.experience_bullets = [
            f"Tailored bullet point #{i} demonstrating leadership and systems programming."
            for i in range(1, 7)
        ]
        self.tools_lines = [
            "Linux, Docker, Git, Bash",
            "WSL, macOS, VS Code",
            "Claude, ChatGPT, Gemini",
            "AntiGravity, Copilot",
        ]
        self.core_concepts_lines = [
            "Multithreading, Concurrency, OOP",
            "Networking, Memory Management",
        ]
        self.coursework_line = "Core CS Coursework: Systems Programming, Operating Systems, Computer Architecture (91)."
        self.military_bullets = [
            "Commanded squads in high-pressure operational environments.",
            "Directed logistics and strategic resource coordination.",
            "Led tactical logistics team ensuring operational readiness.",
        ]


@pytest.mark.skipif(not TEMPLATE_PATH.exists(), reason="Template PPTX not available")
def test_reorder_pill_groups():
    prs = Presentation(TEMPLATE_PATH)
    slide = prs.slides[0]
    group_51 = None
    for shape in slide.shapes:
        if shape.name == "Group 51":
            group_51 = shape
            break

    assert group_51 is not None, "Group 51 must exist in template"
    # Should run cleanly without exception
    reorder_pill_groups(group_51, ["Python", "C++", "SQL", "Java", "TypeScript", "C", "Assembly x86"])


@pytest.mark.skipif(not TEMPLATE_PATH.exists(), reason="Template PPTX not available")
def test_apply_full_tailoring():
    prs = Presentation(TEMPLATE_PATH)
    tailored = DummyTailored()
    apply_full_tailoring(prs, tailored)

    slide = prs.slides[0]
    # Check that TextBox 43 has updated projects
    for shape in slide.shapes:
        if shape.name == "TextBox 43":
            assert "Sockets, Reactor" in shape.text_frame.text
        elif shape.name == "TextBox 36":
            assert "Tailored bullet point #1" in shape.text_frame.text
        elif shape.name == "TextBox 34":
            assert "Operating Systems" in shape.text_frame.text
